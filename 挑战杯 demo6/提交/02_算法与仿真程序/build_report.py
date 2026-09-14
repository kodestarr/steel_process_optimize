"""Build a production-focused scheduling report for one delivery mode."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


COLORS = {
    "text": "0F172A",
    "secondary": "475569",
    "muted": "94A3B8",
    "primary": "3B82F6",
    "primary_dark": "1D4ED8",
    "primary_light": "EFF6FF",
    "success": "10B981",
    "success_light": "ECFDF5",
    "warning": "F59E0B",
    "warning_light": "FFFBEB",
    "danger": "EF4444",
    "danger_light": "FEF2F2",
    "surface": "F7F8FA",
    "white": "FFFFFF",
    "border": "E2E8F0",
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _set_font(run, size: float, color: str = COLORS["text"], bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.bold = bold


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_border(cell, color: str = COLORS["border"], size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(width_cm * 567)))
    width.set(qn("w:type"), "dxa")


def _disable_autofit(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_width(table, width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(int(width_cm * 567)))
    width.set(qn("w:type"), "dxa")


def add_paragraph(
    doc: Document,
    text: str,
    *,
    size: float = 10.5,
    color: str = COLORS["text"],
    bold: bool = False,
    before: float = 0,
    after: float = 7,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.28
    run = paragraph.add_run(text)
    _set_font(run, size, color, bold)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _set_font(run, 14, "000000", True)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    header_fill: str = COLORS["primary_dark"],
    font_size: float = 9.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    _disable_autofit(table)
    _set_table_width(table, sum(widths))
    header = table.rows[0]
    _set_row_repeat(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        _set_cell_width(cell, widths[index])
        _set_cell_shading(cell, header_fill)
        _set_cell_border(cell)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(value))
        _set_font(run, font_size, COLORS["white"], True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            _set_cell_width(cell, widths[index])
            _set_cell_border(cell)
            _set_cell_margins(cell)
            if row_index % 2 == 1:
                _set_cell_shading(cell, "F8FAFC")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(str(value))
            _set_font(run, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def number(value, digits: int = 2, default: str = "N/A") -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return default
    if pd.isna(numeric):
        return default
    return f"{numeric:.{digits}f}"


def percent(value, digits: int = 1) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "N/A"
    if pd.isna(numeric):
        return "N/A"
    if abs(numeric) <= 1.5:
        numeric *= 100
    return f"{numeric:.{digits}f}%"


def add_image(doc: Document, path: Path, caption: str, width: float = 17.2) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    add_paragraph(
        doc,
        caption,
        size=8.8,
        color=COLORS["secondary"],
        after=8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def add_metric_cards(doc: Document, metrics: dict, base: dict) -> None:
    definitions = [
        ("总完工时间", "总完工时间(h)", "h", "lower"),
        ("加权平均齐套跨度", "加权平均齐套跨度(h)", "h", "lower"),
        ("整体产能", "整体产能(张板/班)", "张/班", "higher"),
        ("切割机综合稼动率", "切割机综合稼动率", "%", "higher"),
        ("切割负载差", "切割负载差(h)", "h", "lower"),
        ("死锁与冲突", "风险", "项", "lower"),
    ]
    table = doc.add_table(rows=2, cols=3)
    _disable_autofit(table)
    _set_table_width(table, 17.4)
    risks = int(metrics.get("死锁警告数", 0) or 0) + int(metrics.get("AGV路径冲突次数", 0) or 0)
    values = {
        "总完工时间(h)": metrics.get("总完工时间(h)"),
        "加权平均齐套跨度(h)": metrics.get("加权平均齐套跨度(h)"),
        "整体产能(张板/班)": metrics.get("整体产能(张板/班)"),
        "切割机综合稼动率": metrics.get("切割机综合稼动率"),
        "切割负载差(h)": metrics.get("切割负载差(h)"),
        "风险": risks,
    }
    for index, (label, key, unit, better) in enumerate(definitions):
        row = index // 3
        column = index % 3
        cell = table.cell(row, column)
        _set_cell_width(cell, 5.8)
        _set_cell_border(cell)
        _set_cell_margins(cell, 140, 160, 140, 160)
        _set_cell_shading(cell, COLORS["surface"])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        label_run = paragraph.add_run(label)
        _set_font(label_run, 8.5, COLORS["secondary"], False)
        value_paragraph = cell.add_paragraph()
        value_paragraph.paragraph_format.space_before = Pt(3)
        value_paragraph.paragraph_format.space_after = Pt(2)
        value = values[key]
        if key == "切割机综合稼动率":
            display = percent(value)
            suffix = ""
        elif key == "风险":
            display = str(risks)
            suffix = f" {unit}"
        else:
            display = number(value)
            suffix = f" {unit}"
        value_run = value_paragraph.add_run(f"{display}{suffix}")
        accent = COLORS["primary"] if better == "higher" else COLORS["text"]
        if key == "风险" and risks > 0:
            accent = COLORS["danger"]
        _set_font(value_run, 16, accent, True)
        baseline = base.get(key)
        if baseline is not None and key != "风险":
            note = cell.add_paragraph()
            note.paragraph_format.space_after = Pt(0)
            baseline_display = (
                percent(baseline)
                if key == "切割机综合稼动率"
                else number(baseline)
            )
            note_run = note.add_run(f"FIFO {baseline_display}")
            _set_font(note_run, 8, COLORS["secondary"], False)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def load_inputs(results: Path):
    summary = json.loads((results / "summary.json").read_text(encoding="utf-8"))
    comparison = pd.read_csv(results / "comparison.csv")
    schedule = pd.read_csv(results / "optimized_plate_schedule.csv")
    groups = pd.read_csv(results / "kit_groups.csv")
    stages = pd.read_csv(results / "process_stages.csv")
    return summary, comparison, schedule, groups, stages


def build(results: Path, output: Path, mode: str, run_id: str) -> None:
    summary, comparison, schedule, groups, stages = load_inputs(results)
    base = summary.get("FIFO基线", {})
    optimized = summary.get("齐套感知优化", {})
    checks = summary.get("数据校验", {})
    algorithm = str(summary.get("算法名称") or "优化调度")
    mode_label = "产能优先方案" if mode == "capacity" else "兼顾三指标方案"

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("船体加工车间智能排产与齐套配盘优化调度生产报告")
    _set_font(run, 20, "000000", True)
    add_paragraph(
        doc,
        f"{mode_label}  |  运行ID {run_id}  |  生成时间 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        size=9.3,
        color=COLORS["secondary"],
        after=10,
    )

    add_heading(doc, "生产结论")
    cmax = optimized.get("总完工时间(h)")
    kit = optimized.get("加权平均齐套跨度(h)")
    throughput = optimized.get("整体产能(张板/班)")
    oee = optimized.get("切割机综合稼动率")
    add_paragraph(
        doc,
        f"建议按本报告的{mode_label}执行。优化总完工时间为 {number(cmax)} h，"
        f"加权平均齐套跨度为 {number(kit)} h，整体产能为 {number(throughput)} 张/班，"
        f"切割机综合稼动率为 {percent(oee)}。",
        size=10.5,
        after=8,
    )
    add_paragraph(
        doc,
        f"结果来源：{algorithm}。完整明细见本目录中的“排产结果”文件夹，报告仅保留生产决策和现场核对所需信息。",
        size=9.3,
        color=COLORS["secondary"],
        after=8,
    )

    add_metric_cards(doc, optimized, base)
    add_heading(doc, "基线与优化方案对比")
    comparison_rows = []
    for _, row in comparison.iterrows():
        comparison_rows.append([
            str(row.get("方案", "")),
            number(row.get("总完工时间(h)")),
            number(row.get("加权平均齐套跨度(h)")),
            number(row.get("最大齐套跨度(h)")),
            number(row.get("整体产能(张板/班)")),
            number(row.get("切割负载差(h)")),
        ])
    add_table(
        doc,
        ["方案", "总完工(h)", "平均齐套跨度(h)", "最大齐套跨度(h)", "产能(张/班)", "负载差(h)"],
        comparison_rows,
        [3.2, 2.6, 3.1, 3.1, 2.8, 2.4],
    )

    doc.add_page_break()
    add_heading(doc, "关键运行图")
    add_image(doc, results / "cutting_gantt.png", "切割机排程甘特图", 17.2)
    figure_table = doc.add_table(rows=1, cols=2)
    _disable_autofit(figure_table)
    _set_table_width(figure_table, 17.4)
    for column, (filename, caption) in enumerate((
        ("kit_span.png", "齐套跨度分布"),
        ("resource_utilisation.png", "设备利用率"),
    )):
        cell = figure_table.cell(0, column)
        _set_cell_width(cell, 8.7)
        _set_cell_border(cell, COLORS["white"], 0)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_path = results / filename
        if image_path.exists():
            paragraph.add_run().add_picture(str(image_path), width=Cm(8.0))
        caption_paragraph = cell.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_paragraph.add_run(caption)
        _set_font(caption_run, 8.5, COLORS["secondary"])

    add_heading(doc, "生产排程摘录")
    schedule_view = schedule.sort_values("切割开始(min)", kind="stable").head(24)
    rows = []
    for _, row in schedule_view.iterrows():
        rows.append([
            str(row.get("切割序号", "")),
            str(row.get("套料图名", "")),
            str(row.get("分段号", "")),
            str(row.get("切割机", "")),
            number(float(row.get("切割开始(min)", 0)) / 60, 2),
            number(float(row.get("切割完成(min)", 0)) / 60, 2),
            number(float(row.get("切割工时(min)", 0)) / 60, 2),
        ])
    add_table(
        doc,
        ["序号", "钢板", "分段", "设备", "开始(h)", "完成(h)", "工时(h)"],
        rows,
        [1.2, 4.4, 2.4, 1.6, 2.0, 2.0, 2.0],
        font_size=7.9,
    )
    add_paragraph(
        doc,
        f"本表列出前 {len(schedule_view)} 张钢板，共 {len(schedule)} 张。完整排程、零件完工时间和全部工序事件见“排产结果”目录。",
        size=8.8,
        color=COLORS["secondary"],
        after=8,
    )

    doc.add_page_break()
    add_heading(doc, "资源负荷与瓶颈")
    resource_rows = []
    for label, key in (
        ("切割机综合稼动率", "切割机综合稼动率"),
        ("切割机纯切割稼动率", "切割机纯切割稼动率"),
        ("N2利用率", "N2利用率"),
        ("N5利用率", "N5利用率"),
        ("人工坡口1利用率", "人工坡口1利用率"),
        ("自动坡口1利用率", "自动坡口1利用率"),
        ("AGV1利用率", "AGV1利用率"),
        ("AGV2利用率", "AGV2利用率"),
    ):
        if key in optimized:
            resource_rows.append([label, percent(optimized.get(key))])
    resource_rows.extend([
        ["坡口缓存峰值", percent(optimized.get("坡口缓存区峰值占用率"))],
        ["半框缓存峰值", percent(optimized.get("半框缓存区峰值占用率"))],
        ["齐套缓存峰值", percent(optimized.get("齐套缓存区峰值占用率"))],
        ["AGV路径冲突", str(optimized.get("AGV路径冲突次数", 0))],
        ["死锁警告", str(optimized.get("死锁警告数", 0))],
    ])
    add_table(doc, ["资源/诊断项", "结果"], resource_rows, [10.0, 7.4])

    doc.add_page_break()
    add_heading(doc, "数据完整性与执行检查")
    add_table(
        doc,
        ["检查项", "结果"],
        [
            ["钢板数", str(checks.get("钢板数", "N/A"))],
            ["零件数", str(checks.get("零件数", "N/A"))],
            ["齐套组数", str(checks.get("齐套组数(分段+优先级)", "N/A"))],
            ["未匹配零件", str(checks.get("未匹配零件数", "N/A"))],
            ["打磨长度缺失", str(checks.get("打磨长度缺失零件数", "N/A"))],
            ["工艺参数来源", str(checks.get("工艺参数来源", summary.get("工艺参数来源", "N/A")))],
        ],
        [8.0, 9.4],
    )

    add_heading(doc, "交付文件")
    add_table(
        doc,
        ["文件", "用途"],
        [
            ["optimized_plate_schedule.csv", "钢板切割顺序、设备分配和时间"],
            ["part_completion.csv", "零件完工与齐套时间"],
            ["process_stages.csv", "全部资源工序事件"],
            ["comparison.csv", "FIFO 基线与当前方案对比"],
            ["cutting_gantt.png", "切割排程甘特图"],
            ["kit_span.png", "齐套跨度分布"],
            ["resource_utilisation.png", "资源利用率"],
        ],
        [6.2, 11.2],
    )

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(f"运行ID {run_id}  |  第 ")
    _set_font(footer_run, 8, COLORS["secondary"])
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer_paragraph._p.append(field)
    end_run = footer_paragraph.add_run(" 页")
    _set_font(end_run, 8, COLORS["secondary"])

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--mode", choices=("capacity", "balanced"), default="capacity")
    parser.add_argument("--run-id", default="")
    args = parser.parse_args()
    run_id = args.run_id or args.results.parent.parent.name.rsplit("_", 1)[-1]
    build(args.results, args.output, args.mode, run_id)
    print(args.output)


if __name__ == "__main__":
    main()
